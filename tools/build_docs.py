# -*- coding: utf-8 -*-
"""
build_docs.py — 原稿 (content_*.py) と実装ファイル・検証結果から
  JBossEAP81_WARN13_解説.xlsx (Meiryo UI / モノトーン) と JBossEAP81_WARN13_解説.md を生成する。

使い方:  python tools/build_docs.py   (WarnList フォルダで実行)
必要:    openpyxl
"""
import os
import re
import sys

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)

import content_meta as M  # noqa: E402
from content_deep import DEEP1  # noqa: E402
from content_w01_04 import W01, W02, W03, W04  # noqa: E402
from content_w05_09 import W05, W06, W07, W08, W09  # noqa: E402
from content_w10_13 import W10, W11, W12, W13  # noqa: E402
from doclib import (BAND_FILL, CODE_FILL, FONT, HEAD_FILL, INK, INK_SUB, LINE, SUB_FILL, WHITE, C, P, S, Sheet, T,  # noqa: E402
                    count_lines, md_blocks, md_sections, new_workbook)

WARNINGS = [W01, W02, W03, W04, W05, W06, W07, W08, W09, W10, W11, W12, W13]
RES = os.path.join(ROOT, "verification", "results")
IMPL = os.path.join(ROOT, "implementation")
OUT_XLSX = os.path.join(ROOT, "JBossEAP81_WARN13_解説.xlsx")
OUT_MD = os.path.join(ROOT, "JBossEAP81_WARN13_解説.md")

SCENARIOS = [
    ("s0-before", "S0 修正前", "EAP 既定 + ExtraSLB 方式の信頼設定 + 修正前 WAR + 修正前データソース。SIGTERM は JVM へ中継 (LAUNCH_JBOSS_IN_BACKGROUND=true)、サスペンド待ち 0 秒"),
    ("s0n-before-nolaunch", "S0n 修正前・中継なし", "S0 と同じで LAUNCH_JBOSS_IN_BACKGROUND 未設定 (SIGTERM が JVM に届かない場合)"),
    ("s1-after-system", "S1 修正後 (system/remove)", "本書の実装: TRUST_MODE=system、INBOUND_HTTPS_MODE=remove + 修正後 WAR。EAP_SIGTERM_SUSPEND_TIMEOUT=30、DB_IDLE_TIMEOUT_MIN=1、DB_TRACK_STATEMENTS=TRUE"),
    ("s1e-after-elytron", "S1e 修正後 (elytron)", "S1 の TRUST_MODE を elytron (ExtraSLB 方式 + required=true) にしたもの"),
    ("s1d-after-dedicated", "S1d 修正後 (dedicated)", "S1 の INBOUND_HTTPS_MODE を dedicated (起動時に渡す正式な鍵ストア) にしたもの"),
]

LINK_COLOR = "404040"


# =====================================================================
# 検証結果の読み込み
# =====================================================================
def read(path):
    with open(path, encoding="utf-8", errors="replace") as f:
        return f.read()


def load_counts():
    counts = {}
    for key, _, _ in SCENARIOS:
        p = os.path.join(RES, f"{key}-warnings.txt")
        d = {}
        if os.path.exists(p):
            for line in read(p).splitlines():
                m = re.search(r"(\d+)\s*$", line)
                if m and not line.startswith(("No", "---", "total")):
                    d[line.strip()] = int(m.group(1))
        counts[key] = d
    return counts


COUNTS = load_counts()


def count_of(scenario, label):
    for line, n in COUNTS.get(scenario, {}).items():
        if label.strip() in line:
            return n
    return None


def verify_section(w):
    cols = [s for s in SCENARIOS if s[0] != "s0n-before-nolaunch" or w["key"] == "W13"]
    header = ["集計項目 (check-warnings.sh)"] + [s[1] for s in cols]
    rows = []
    for lb in w["verify_labels"]:
        rows.append([lb.strip()] + [("-" if count_of(s[0], lb) is None else str(count_of(s[0], lb))) for s in cols])
    blocks = [T(header, rows, widths=[4] + [2] * len(cols)),
              P("S0 = 修正前を再現した状態、S1 系 = 本書の実装を適用した状態。シナリオの詳細と全ログは「検証結果」シート (md では「検証結果」章)。")]
    return S("検証結果 (実測)", *blocks)


def refs_rows(keys):
    return [[M.REF_BY_KEY[k][1], M.REF_BY_KEY[k][2]] for k in keys]


# =====================================================================
# 実装ファイル
# =====================================================================
IMPL_MAP = [
    ("container/Dockerfile", "2 段ビルド (UBI 9.8 でEAP 展開・CLI 適用 → ubi9/openjdk-21-runtime で実行)。信頼設定・ドライバ・CLI をまとめる", "①③④⑥⑧⑨⑩⑪⑫⑬", "ビルド"),
    ("container/cli/00-system-properties.cli", "SIGTERM 時のサスペンド待ち時間 (org.wildfly.sigterm.suspend.timeout)", "⑬", "ビルド (embed-server)"),
    ("container/cli/10-elytron-outbound-truststore.cli", "elytron モードの信頼設定 (ExtraSLB と同じ構成 + required=true)", "①", "ビルド (TRUST_MODE=elytron)"),
    ("container/cli/11-inbound-https-remove.cli", "使っていないインバウンド HTTPS 一式 (8443) を削除", "①", "ビルド (INBOUND_HTTPS_MODE=remove)"),
    ("container/cli/12-inbound-https-dedicated.cli", "8443 を起動時に渡す正式な鍵ストアで動かす", "①", "ビルド (INBOUND_HTTPS_MODE=dedicated)"),
    ("container/cli/20-logging.cli", "④ Spring 分の WFLYEE0007 と ③ spring-core 由来の org.jboss.vfs の WFLYSRV0018 だけを抑止。調査用コマンド", "③④⑥", "ビルド (embed-server)"),
    ("container/cli/30-datasource-mysql-xa.cli", "JDBC ドライバと XA データソース (検証・exception-sorter・idle-timeout・track-statements・TCP keepalive)", "⑧⑩⑪⑫", "ビルド (embed-server)"),
    ("container/cli/31-jca-ccm.cli", "接続リーク検出 (CCM) を環境変数で切り替え", "⑧⑫⑬", "ビルド (embed-server)"),
    ("container/cli/60-undertow.cli", "X-Forwarded-* の採用と 4xx/5xx の JSON アクセスログ", "⑨", "ビルド (embed-server)"),
    ("container/cli/70-transactions.cli", "XA の node-identifier をタスクごとに一意に", "(XA の前提)", "ビルド (embed-server)"),
    ("container/cli/90-verify.cli", "適用結果の表示", "全般", "ビルド / 起動後"),
    ("container/scripts/apply-cli.sh", "embed-server で cli/ を番号順に適用し、standalone_xml_history を残さない", "全般", "ビルド"),
    ("container/scripts/install-ca-anchors.sh", "受領した CA を UBI のシステムトラストストアへ追加 (update-ca-trust) し、cacerts 入りを指紋で確認", "①", "ビルド (TRUST_MODE=system)"),
    ("container/scripts/build-truststore.sh", "Container_ExtraSLB_JVM_https_outbounds の同名スクリプト (変更なし)。cacerts のコピー + CA の p12", "①", "ビルド (TRUST_MODE=elytron)"),
    ("container/scripts/entrypoint.sh", "起動前チェック、TX_NODE_ID、SIGTERM の中継 (LAUNCH_JBOSS_IN_BACKGROUND)、standalone.sh 起動", "①⑬", "起動"),
    ("container/modules/com/mysql/main/module.xml", "MySQL Connector/J のモジュール定義 (デプロイではなくモジュールで)", "⑩⑪⑬", "ビルド"),
    ("app/src/main/webapp/WEB-INF/jboss-deployment-structure.xml", "修正後の記述子 (testClasses・cxf.impl・誤った除外を削除、jsf と weld を除外)", "②③⑤⑥", "WAR"),
    ("app/pom-snippets.xml", "サーバ提供 API を provided に、javax.* と log4j-core の混入を enforcer で禁止", "②⑤⑥⑦", "WAR"),
    ("app/src/main/java/com/example/iwinmichl/web/MethodNotAllowedAdvice.java", "405 を送信元・URL 付きの 1 行で記録", "⑨", "WAR"),
    ("app/src/main/java/com/example/iwinmichl/dao/JdbcResourceSamples.java", "JDBC リソースの正しい閉じ方", "⑧⑫", "WAR"),
    ("app/src/main/java/com/example/iwinmichl/ws/SoapClientTimeouts.java", "標準 API だけで SOAP のタイムアウトを設定 (cxf.impl 不要)", "③", "WAR"),
    ("ecs/taskdef-container-fragment.json", "ECS タスク定義 (secrets、stopTimeout、systemControls、healthCheck、awslogs)", "⑩⑪⑬", "デプロイ"),
    ("tools/check-warnings.sh", "ログから警告 ①〜⑬ と関連メッセージの件数を集計", "全般", "運用"),
    ("tools/analyze-ij-log.py", "⑧⑫ の作成箇所、⑪ の放置時間と idle-timeout の推奨値、⑩ の根本原因、⑬ の発生状況", "⑧⑩⑪⑫⑬", "運用"),
    ("tools/scan-war.py", "WAR を静的に検査し ②③④⑤⑥⑦ の原因を特定", "②③④⑤⑥⑦", "ビルド / CI"),
    ("tools/cacert-secret-args.sh", "Container_ExtraSLB_JVM_https_outbounds の同名スクリプト (変更なし)。証明書を 1 つのシークレットに", "①", "ビルド"),
]
UNCHANGED_FROM_EXTRASLB = {"container/scripts/build-truststore.sh", "tools/cacert-secret-args.sh"}


def lang_of(path):
    ext = os.path.splitext(path)[1]
    return {".cli": "text", ".sh": "bash", ".py": "python", ".java": "java", ".xml": "xml", ".json": "json"}.get(ext, "dockerfile")


def impl_files():
    for rel, role, warns, when in IMPL_MAP:
        if rel in UNCHANGED_FROM_EXTRASLB:
            continue
        yield rel, read(os.path.join(IMPL, rel))


# =====================================================================
# Excel: 表 (グリッド) 型シート
# =====================================================================
def grid_sheet(wb, title, header, rows, widths, tab="BFBFBF", sev_col=None, link_cols=(), title_text=None, note=None):
    ws = wb.create_sheet(title=title)
    ws.sheet_properties.tabColor = tab
    ws.sheet_view.showGridLines = False
    thin = Side(style="thin", color=LINE)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    r0 = 1
    if title_text:
        ws.cell(row=1, column=1, value=title_text).font = Font(name=FONT, size=14, bold=True, color=WHITE)
        for c in range(1, len(header) + 1):
            ws.cell(row=1, column=c).fill = PatternFill("solid", fgColor=HEAD_FILL)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(header))
        ws.cell(row=1, column=1).alignment = Alignment(vertical="center")
        ws.row_dimensions[1].height = 30
        r0 = 2
        if note:
            ws.cell(row=2, column=1, value=note).font = Font(name=FONT, size=9, color=INK_SUB)
            ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(header))
            ws.cell(row=2, column=1).alignment = Alignment(wrap_text=True, vertical="center")
            ws.row_dimensions[2].height = max(18, count_lines(note, sum(widths)) * 15)
            r0 = 3
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c, h in enumerate(header, start=1):
        cell = ws.cell(row=r0, column=c, value=h)
        cell.font = Font(name=FONT, size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor="404040")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[r0].height = 30
    for i, row in enumerate(rows):
        r = r0 + 1 + i
        lines = 1
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name=FONT, size=10, color=INK)
            cell.alignment = Alignment(vertical="top", wrap_text=True,
                                       horizontal="center" if (sev_col and c == sev_col) or widths[c - 1] <= 6 else "left")
            cell.border = border
            if i % 2 == 1:
                cell.fill = PatternFill("solid", fgColor=BAND_FILL)
            if sev_col and c == sev_col:
                fill, color = {"A": ("000000", WHITE), "B": ("595959", WHITE), "C": ("D9D9D9", INK)}.get(str(v), (None, INK))
                if fill:
                    cell.fill = PatternFill("solid", fgColor=fill)
                    cell.font = Font(name=FONT, size=10, bold=True, color=color)
            if c in link_cols and isinstance(v, str) and v.startswith("http"):
                cell.hyperlink = v
                cell.font = Font(name=FONT, size=9, color=LINK_COLOR, underline="single")
            lines = max(lines, count_lines(str(v), widths[c - 1]))
        ws.row_dimensions[r].height = min(400, lines * 15 + 5)
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    ws.print_title_rows = f"{r0}:{r0}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def style_links(ws, col_letter="C"):
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("http"):
                cell.hyperlink = cell.value
                cell.font = Font(name=FONT, size=9, color=LINK_COLOR, underline="single")


# =====================================================================
# Excel 全体
# =====================================================================
def build_xlsx():
    wb = new_workbook()
    sheet_names = []

    # ---- 表紙 ----
    cv = Sheet(wb, "表紙", tab_gray="000000")
    cv.title(M.TITLE, M.SUBTITLE)
    cv.heading("対象環境")
    cv.table(["項目", "内容"], M.ENV_ROWS, widths=[1, 5])
    cv.blank()
    cv.heading("危険度の見方")
    cv.table(["危険度", "意味"], M.LEGEND_ROWS, widths=[1, 6])
    cv.blank()
    cv.heading("この資料の読み方")
    cv.labeled("", "\n".join("・" + x for x in M.HOW_TO_READ))
    cv.blank()
    cv.heading("シート一覧 (クリックで移動)")
    index_rows = [["一覧", "13 件の要約と対処・検証結果の一覧"], ["基礎知識", "EAP・JBoss Modules・Elytron・IronJacamar・ECS の停止・ログの読み方・jboss-cli"]]
    index_rows += [[w["sheet"], w["title"]] for w in WARNINGS]
    index_rows += [["①深掘り", "追加質問 Q1〜Q3 の回答と Container_ExtraSLB_JVM_https_outbounds の実装レビュー"],
                   ["実装マップ", "実装ファイルと対応する警告"], ["実装コード", "jboss-cli・Dockerfile・entrypoint・ECS・WAR 側・ツールの全文"],
                   ["検証結果", "修正前の再現と修正後の解消の実測記録"], ["年表", "関係する製品・仕様の歴史"], ["用語集", "小学生向けのひとこと付き"],
                   ["前提・版", "前提にした環境・版・推定した点・確認してほしい点"], ["出典", "根拠にした資料"]]
    start = cv.row
    cv.table(["シート", "内容"], index_rows, widths=[1, 5])
    for i, (name, _) in enumerate(index_rows):
        cell = cv.ws.cell(row=start + 1 + i, column=2)
        cell.hyperlink = f"#'{name}'!A1"
        cell.font = Font(name=FONT, size=10, color=LINK_COLOR, underline="single")
    sheet_names.append("表紙")

    # ---- 一覧 ----
    rows = []
    for r in M.SUMMARY:
        w = next(x for x in WARNINGS if x["key"] == r[-1])
        lb = w["verify_labels"][0]
        rows.append(r[:-1] + [f"{count_of('s0-before', lb)} → {count_of('s1-after-system', lb)}"])
    grid_sheet(wb, "一覧", M.SUMMARY_HEADER, rows, M.SUMMARY_WIDTHS, tab="262626", sev_col=7,
               title_text="一覧: JBoss EAP 8.1 警告 13 件",
               note="検証列は「修正前 (S0) の件数 → 修正後 (S1) の件数」。UBI 9.8 + OpenJDK 21 + WildFly 35.0.1 (EAP 8.1 の上流) + MySQL 8.4 のコンテナで実測。")

    # ---- 基礎知識 ----
    bs = Sheet(wb, "基礎知識", tab_gray="404040")
    bs.title("基礎知識 (先に読むと各警告の説明が分かりやすくなります)")
    bs.sections(M.BASICS)

    # ---- 各警告 ----
    tabs = {"A": "000000", "B": "595959", "C": "A6A6A6"}
    for w in WARNINGS:
        sh = Sheet(wb, w["sheet"], tab_gray=tabs[w["severity"]])
        sh.title(w["title"], f"危険度 {w['severity']}  /  {dict((x[0], x[1]) for x in M.LEGEND_ROWS)[w['severity']]}")
        sh.heading("ユーザー環境のログ (原文)")
        sh.labeled("ログ", w["user_log"])
        sh.blank()
        sh.sections(w["sections"])
        sh.sections([verify_section(w)])
        sh.heading("出典")
        start = sh.row
        sh.table(["資料", "URL"], refs_rows(w["refs"]), widths=[2, 3])
        style_links(sh.ws)

    # ---- ①深掘り ----
    dp = Sheet(wb, "①深掘り", tab_gray="7F7F7F")
    dp.title("①深掘り: 追加質問 Q1〜Q3 と Container_ExtraSLB_JVM_https_outbounds の実装レビュー")
    dp.sections(DEEP1)

    # ---- 実装マップ ----
    grid_sheet(wb, "実装マップ", ["ファイル (implementation/ 配下)", "役割", "対応する警告", "使う時点"],
               [list(x) for x in IMPL_MAP], [52, 70, 16, 24], tab="BFBFBF",
               title_text="実装マップ: implementation/ フォルダのファイルと対応する警告",
               note="ビルド: docker build (BuildKit) / 起動: ECS タスク起動時 / WAR: アプリ側の修正 / 運用: ログ調査。build-truststore.sh と cacert-secret-args.sh は ExtraSLB プロジェクトのものをそのまま同梱。")

    # ---- 実装コード ----
    ic = Sheet(wb, "実装コード", tab_gray="BFBFBF")
    ic.title("実装コード (implementation/ の全文)", "jboss-cli の .cli はビルド時に apply-cli.sh が embed-server で番号順に適用する。${env.X:既定値} は起動時に ECS の環境変数で解決される。")
    for rel, text in impl_files():
        ic.heading("implementation/" + rel)
        ic.code(text)
        ic.blank()

    # ---- 検証結果 ----
    vr = Sheet(wb, "検証結果", tab_gray="7F7F7F")
    vr.title("検証結果: 修正前に 13 件を再現し、修正後に解消することを実測",
             "Docker Desktop 29.8.1 上のコンテナ。EAP 8.1 のバイナリの代わりに上流の WildFly 35.0.1.Final を使用 (実装の Dockerfile に WildFly の配布物を渡してそのままビルド)。")
    vr.sections(verification_sections())

    # ---- 年表・用語集・前提・出典 ----
    grid_sheet(wb, "年表", M.TIMELINE_HEADER, M.TIMELINE, [14, 100, 18], tab="D9D9D9", title_text="年表: 13 件に関係する製品・仕様の歴史")
    grid_sheet(wb, "用語集", M.GLOSSARY_HEADER, M.GLOSSARY, [30, 40, 80], tab="D9D9D9", title_text="用語集")
    pv = Sheet(wb, "前提・版", tab_gray="D9D9D9")
    pv.title("前提・版・推定した点・確認してほしい点")
    pv.sections(versions_sections())
    grid_sheet(wb, "出典", M.REFS_HEADER, M.REFS, [10, 60, 100], tab="D9D9D9", link_cols=(3,), title_text="出典 (2026-10-11 時点で確認)")

    wb.active = 0
    wb.save(OUT_XLSX)
    return OUT_XLSX


def versions_sections():
    return [
        S("前提にした環境と、検証に使った版", T(M.VERSIONS_HEADER, M.VERSIONS, widths=[2, 4, 5])),
        S("推定した点 (実物で確認できなかったこと)", *[P("・" + x) for x in M.ASSUMPTIONS]),
        S("確認してほしい点", *[P("・" + x) for x in M.TO_CONFIRM]),
    ]


def verification_sections():
    secs = []
    secs.append(S("シナリオ", T(["記号", "内容"], [[s[1], s[2]] for s in SCENARIOS], widths=[1, 6]),
                  P("共通: MySQL 8.4.11 (wait_timeout=150 秒)、自己署名ルート CA の HTTPS テストサーバ (extraslb-test)、SOAP バックエンド (ws-backend.war)。"
                    "各シナリオは ブート → 機能確認 → 170 秒放置して DB 接続を切らせて再利用 → 15 秒かかる DB 処理の最中に docker stop (SIGTERM) の順。")))
    labels = []
    for line in read(os.path.join(RES, "s0-before-warnings.txt")).splitlines():
        if re.search(r"\d+\s*$", line) and not line.startswith(("No", "---", "total")):
            labels.append(re.sub(r"\s+\d+\s*$", "", line).strip())
    rows = []
    for lb in labels:
        rows.append([lb] + [str(count_of(s[0], re.sub(r"^\S+\s+", "", lb) if not lb.startswith(("①", "②", "③", "④", "⑤", "⑥", "⑦", "⑧", "⑨", "⑩", "⑪", "⑫", "⑬", "他")) else lb[1:].strip())) for s in SCENARIOS])
    totals = []
    for s in SCENARIOS:
        t = read(os.path.join(RES, f"{s[0]}-warnings.txt")).strip().splitlines()[-1]
        totals.append(re.sub(r".*=\s*", "", t))
    rows.append(["合計 (①〜⑬ 本体)"] + totals)
    main = [s for s in SCENARIOS if s[0] in ("s0-before", "s0n-before-nolaunch", "s1-after-system")]
    idx = [i for i, s in enumerate(SCENARIOS) if s in main]
    rows3 = [[r[0]] + [r[1 + i] for i in idx] for r in rows]
    same = all(COUNTS[k] == COUNTS["s1-after-system"] for k in ("s1e-after-elytron", "s1d-after-dedicated"))
    secs.append(S("警告の件数 (check-warnings.sh)", T(["項目"] + [s[1] for s in main], rows3, widths=[3, 1, 1, 1]),
                  P("S1e (TRUST_MODE=elytron) と S1d (INBOUND_HTTPS_MODE=dedicated) の件数は、全項目で S1 と同じ値でした (①〜⑬ の本体合計 0)。"
                    if same else "S1e / S1d の件数は results/ の *-warnings.txt を参照。")))
    for key, name in (("s0-before", "S0 修正前"), ("s1-after-system", "S1 修正後")):
        secs.append(S(f"機能確認の記録 ({name})", C(read(os.path.join(RES, f"{key}-http.txt")), "text")))
    secs.append(S("SIGTERM が JVM に届かない場合 (S0n)", C(read(os.path.join(RES, "s0n-before-nolaunch-http.txt")).split("### ⑬")[-1].strip(), "text"),
                  P("standalone.sh が PID 1 で LAUNCH_JBOSS_IN_BACKGROUND が未設定だと、SIGTERM が JVM に届かず、docker stop の猶予 (90 秒) を使い切って SIGKILL (終了コード 137) された。"
                    "停止ログも出ない (処理中の要求は猶予の間に終わったので成功している)。")))
    secs.append(S("① の追加実験 (Q1・Q2・方式比較)", C(read(os.path.join(RES, "tls-experiments.txt")), "text")))
    secs.append(S("④ の原因 (DEBUG) と ⑥ の除外指定の実験", C(read(os.path.join(RES, "extra-experiments.txt")), "text")))
    secs.append(S("修正前ログの抜粋 (S0)", C(s0_excerpt(), "text")))
    secs.append(S("IronJacamar ログの解析 (S0)", C(run_analyzer(), "text")))
    secs.append(S("再現手順", C("""# 1. 部品の取得 (WildFly 35.0.1 / Spring 6.2.19 / jstl-1.2 / jakarta.transaction-api / Connector/J 26.7.0 等を Maven Central から)
# 2. テスト用 PKI と WAR
bash verification/infra/make-test-pki.sh <pki>
bash verification/build-test-wars.sh <jars> <wars>
# 3. イメージ
DOCKER_BUILDKIT=1 docker build --build-context dist=<dl> --build-context impl=implementation/container \\
  --secret id=cacerts,src=<pki>/cacerts-bundle.tar -t wl-before:1 verification/before
DOCKER_BUILDKIT=1 docker build --build-context dist=<dl> --secret id=cacerts,src=<pki>/cacerts-bundle.tar \\
  --build-arg EAP_DIST=wildfly-dist-35.0.1.Final.tar.gz --build-arg TRUST_MODE=system --build-arg INBOUND_HTTPS_MODE=remove \\
  -t wl-after-system-remove:1 implementation/container
# 4. MySQL 8.4 (wait_timeout=150) と HTTPS テストサーバを wl-net 上に起動
# 5. シナリオ
WARS=<wars> PKI=<pki> bash verification/run-all.sh
WARS=<wars> L4J_JARS=<log4j jars> JARS=<jars> bash verification/run-extra-experiments.sh""", "bash")))
    return secs


def s0_excerpt():
    log = read(os.path.join(RES, "s0-before-server.log")).splitlines()
    pats = ["WFLYELY00023", "WFLYELY01084", "WFLYSRV0058", "WFLYSRV0018", "Failed to define class javax", "WFLYEE0007",
            "WFLYWELD0013", "WFLYSRV0274", "IJ030040", "IJ030042", "Resolved \\[", "IJ030027", "Caused by: com.mysql",
            "IJ000621", "WFLYSRV0272", "WFLYSRV0220", "IJ000615", "IJ030022"]
    out = []
    for p in pats:
        for line in log:
            if re.search(p, line):
                out.append(line[:420])
                break
    return "\n".join(out)


def run_analyzer():
    import subprocess
    try:
        r = subprocess.run([sys.executable, os.path.join(IMPL, "tools", "analyze-ij-log.py"),
                            os.path.join(RES, "s0-before-server.log"), "--pkg", "com.example"],
                           capture_output=True, text=True, encoding="utf-8")
        return r.stdout.strip()
    except Exception as e:  # pragma: no cover
        return f"(実行できませんでした: {e})"


# =====================================================================
# Markdown
# =====================================================================
def build_md():
    out = [f"# {M.TITLE}\n", f"{M.SUBTITLE}\n"]
    toc = ["## 目次\n", "- [対象環境と読み方](#対象環境と読み方)", "- [一覧](#一覧)", "- [基礎知識](#基礎知識)"]
    for w in WARNINGS:
        toc.append(f"- [{w['title']}](#{anchor(w['title'])})")
    toc += ["- [①深掘り: 追加質問と ExtraSLB 実装レビュー](#①深掘り-追加質問と-extraslb-実装レビュー)", "- [実装](#実装)", "- [検証結果](#検証結果)",
            "- [年表](#年表)", "- [用語集](#用語集)", "- [前提・版](#前提版)", "- [出典](#出典)", ""]
    out += toc
    out.append("## 対象環境と読み方\n")
    out += md_blocks([T(["項目", "内容"], M.ENV_ROWS), T(["危険度", "意味"], M.LEGEND_ROWS)])
    out += ["- " + x for x in M.HOW_TO_READ] + [""]
    out.append("## 一覧\n")
    rows = []
    for r in M.SUMMARY:
        w = next(x for x in WARNINGS if x["key"] == r[-1])
        lb = w["verify_labels"][0]
        rows.append(r[:-1] + [f"{count_of('s0-before', lb)} → {count_of('s1-after-system', lb)}"])
    out += md_blocks([T(M.SUMMARY_HEADER, rows)])
    out.append("## 基礎知識\n")
    out += md_sections(M.BASICS, level=3)
    for w in WARNINGS:
        out.append(f"## {w['title']}\n")
        out.append(f"**危険度 {w['severity']}**\n")
        out.append("**ユーザー環境のログ (原文)**\n")
        out.append("```text\n" + w["user_log"] + "\n```\n")
        out += md_sections(w["sections"], level=3)
        out += md_sections([verify_section(w)], level=3)
        out.append("### 出典\n")
        out += [f"- [{M.REF_BY_KEY[k][1]}]({M.REF_BY_KEY[k][2]})" for k in w["refs"]] + [""]
    out.append("## ①深掘り: 追加質問と ExtraSLB 実装レビュー\n")
    out += md_sections(DEEP1, level=3)
    out.append("## 実装\n")
    out += md_blocks([T(["ファイル (implementation/ 配下)", "役割", "対応する警告", "使う時点"], [list(x) for x in IMPL_MAP])])
    for rel, text in impl_files():
        out.append(f"### implementation/{rel}\n")
        out.append(f"```{lang_of(rel)}\n{text.rstrip()}\n```\n")
    out.append("## 検証結果\n")
    out += md_sections(verification_sections(), level=3)
    out.append("## 年表\n")
    out += md_blocks([T(M.TIMELINE_HEADER, M.TIMELINE)])
    out.append("## 用語集\n")
    out += md_blocks([T(M.GLOSSARY_HEADER, M.GLOSSARY)])
    out.append("## 前提・版\n")
    out += md_sections(versions_sections(), level=3)
    out.append("## 出典\n")
    out += md_blocks([T(M.REFS_HEADER, [[a, b, f"<{c}>"] for a, b, c in M.REFS])])
    with open(OUT_MD, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(out).rstrip() + "\n")
    return OUT_MD


def anchor(title):
    a = title.strip().lower()
    a = re.sub(r"[^\w\- ぁ-んァ-ヶ一-龠ー①-⑬]", "", a)
    return a.replace(" ", "-")


if __name__ == "__main__":
    x = build_xlsx()
    m = build_md()
    print("wrote", x)
    print("wrote", m)
