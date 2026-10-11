# -*- coding: utf-8 -*-
"""
doclib.py — 1 つの原稿 (Python のデータ) から Excel (Meiryo UI / モノトーン) と Markdown を生成する部品

原稿のブロック:
  P(text)                 段落
  B(*items)               箇条書き
  T(header, rows, widths) 表 (widths は Excel の列幅の比率。省略可)
  C(text, lang, caption)  コード (等幅ではなく Meiryo UI で表示する。背景を薄い灰色にする)
  KID(text)               小学生向けの説明 (囲み)
  NOTE(text)              注意・補足 (囲み)
  DIA(text, caption)      テキスト図 (動作イメージ)
  S(title, *blocks)       見出し付きのまとまり
"""
import math
import re
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

FONT = "Meiryo UI"

# ---- モノトーンの配色 (黒・白・灰色だけ) ----
INK = "000000"        # 本文
INK_SUB = "404040"    # 補助文字
HEAD_FILL = "262626"  # 見出し帯 (濃灰) / 白文字
SUB_FILL = "D9D9D9"   # 小見出し・表の見出し
BAND_FILL = "F2F2F2"  # ラベル列・交互の行
CODE_FILL = "EDEDED"  # コード
KID_FILL = "FAFAFA"   # 小学生向け
LINE = "BFBFBF"       # 罫線
WHITE = "FFFFFF"


def P(text):
    return ("p", text)


def B(*items):
    return ("bullets", list(items))


def T(header, rows, widths=None):
    return ("table", {"header": header, "rows": rows, "widths": widths})


def C(text, lang="", caption=""):
    return ("code", {"text": text.rstrip("\n"), "lang": lang, "caption": caption})


def KID(text):
    return ("kid", text)


def NOTE(text):
    return ("note", text)


def DIA(text, caption=""):
    return ("diagram", {"text": text.rstrip("\n"), "caption": caption})


def S(title, *blocks):
    return {"title": title, "blocks": list(blocks)}


# =====================================================================
# Markdown
# =====================================================================
def md_escape_cell(s):
    s = str(s).replace("|", "\\|")
    return s.replace("\n", "<br>")


def md_blocks(blocks):
    out = []
    for kind, data in blocks:
        if kind == "p":
            out.append(data + "\n")
        elif kind == "bullets":
            out.extend(f"- {x}" for x in data)
            out.append("")
        elif kind == "table":
            hdr = data["header"]
            out.append("| " + " | ".join(md_escape_cell(h) for h in hdr) + " |")
            out.append("|" + "|".join("---" for _ in hdr) + "|")
            for r in data["rows"]:
                cells = list(r) + [""] * (len(hdr) - len(r))
                out.append("| " + " | ".join(md_escape_cell(c) for c in cells) + " |")
            out.append("")
        elif kind == "code":
            if data["caption"]:
                out.append(f"**{data['caption']}**\n")
            out.append(f"```{data['lang']}")
            out.append(data["text"])
            out.append("```\n")
        elif kind == "kid":
            out.append("> **【小学生にもわかる説明】**  ")
            out.extend("> " + line if line else ">" for line in data.split("\n"))
            out.append("")
        elif kind == "note":
            out.append("> **【注意・補足】**  ")
            out.extend("> " + line if line else ">" for line in data.split("\n"))
            out.append("")
        elif kind == "diagram":
            if data["caption"]:
                out.append(f"**{data['caption']}**\n")
            out.append("```text")
            out.append(data["text"])
            out.append("```\n")
    return out


def md_sections(sections, level=3):
    out = []
    for sec in sections:
        out.append("#" * level + " " + sec["title"] + "\n")
        out.extend(md_blocks(sec["blocks"]))
    return out


# =====================================================================
# Excel
# =====================================================================
def disp_width(s):
    """表示幅 (全角=2, 半角=1)。"""
    w = 0
    for ch in s:
        w += 2 if unicodedata.east_asian_width(ch) in ("W", "F", "A") else 1
    return w


def count_lines(text, width_units):
    """セル幅 (Excel の列幅単位) に折り返したときの行数を見積もる。"""
    cap = max(8.0, width_units * 1.0 - 2)
    lines = 0
    for para in str(text).split("\n"):
        w = disp_width(para)
        lines += max(1, math.ceil(w / cap))
    return lines


LINE_PT = 15.0      # Meiryo UI 10pt の 1 行あたりの高さ (余裕込み)
PAD_PT = 5.0
MAX_ROW_PT = 400.0  # Excel の上限 409pt に余裕を持たせる


def chunk_text(text, width_units, max_lines=24):
    """1 セルに入れると高さ上限を超える長文を、段落 (改行) 単位で複数セルに分ける。"""
    chunks, cur, cur_lines = [], [], 0
    for para in str(text).split("\n"):
        pl = count_lines(para, width_units)
        if cur and cur_lines + pl > max_lines:
            chunks.append("\n".join(cur))
            cur, cur_lines = [], 0
        cur.append(para)
        cur_lines += pl
    if cur:
        chunks.append("\n".join(cur))
    return chunks


class Sheet:
    """見出し列 (A) + 本文列 (B〜G) の文書型シート。"""

    def __init__(self, wb, title, tab_gray="808080", label_w=18, body_ws=(22, 22, 22, 22, 22, 30)):
        self.ws = wb.create_sheet(title=title)
        self.ws.sheet_properties.tabColor = tab_gray
        self.ws.sheet_view.showGridLines = False
        self.ws.sheet_view.zoomScale = 100
        self.label_w = label_w
        self.body_ws = list(body_ws)
        self.ncols = 1 + len(self.body_ws)
        self.ws.column_dimensions["A"].width = label_w
        for i, w in enumerate(self.body_ws, start=2):
            self.ws.column_dimensions[get_column_letter(i)].width = w
        self.row = 1
        thin = Side(style="thin", color=LINE)
        self.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        self.page_setup()

    # ---- 共通 ----
    def page_setup(self):
        ws = self.ws
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5)
        ws.print_options.horizontalCentered = True
        ws.oddFooter.center.text = "&P / &N"
        ws.oddFooter.center.font = FONT

    def font(self, bold=False, color=INK, size=10, underline=None, italic=False):
        return Font(name=FONT, size=size, bold=bold, color=color, underline=underline, italic=italic)

    def body_width(self, start=0, end=None):
        end = len(self.body_ws) if end is None else end
        return sum(self.body_ws[start:end])

    def put(self, r, c1, c2, value, *, bold=False, fill=None, color=INK, size=10, align="left",
            valign="top", wrap=True, border=True, italic=False):
        ws = self.ws
        cell = ws.cell(row=r, column=c1, value=value)
        cell.font = self.font(bold=bold, color=color, size=size, italic=italic)
        cell.alignment = Alignment(horizontal=align, vertical=valign, wrap_text=wrap)
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        if c2 > c1:
            ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        for c in range(c1, c2 + 1):
            cc = ws.cell(row=r, column=c)
            if border:
                cc.border = self.border
            if fill and c != c1:
                cc.fill = PatternFill("solid", fgColor=fill)
        return cell

    def height(self, r, lines):
        self.ws.row_dimensions[r].height = min(MAX_ROW_PT, lines * LINE_PT + PAD_PT)

    # ---- 部品 ----
    def title(self, text, sub=None):
        r = self.row
        self.put(r, 1, self.ncols, text, bold=True, fill=HEAD_FILL, color=WHITE, size=14, valign="center")
        self.ws.row_dimensions[r].height = 30
        self.row += 1
        if sub:
            r = self.row
            self.put(r, 1, self.ncols, sub, color=INK_SUB, fill=BAND_FILL, valign="center")
            self.height(r, count_lines(sub, self.label_w + self.body_width()))
            self.row += 1
        self.row += 1

    def heading(self, text):
        r = self.row
        self.put(r, 1, self.ncols, text, bold=True, fill=SUB_FILL, size=11, valign="center")
        self.ws.row_dimensions[r].height = 22
        self.row += 1

    def labeled(self, label, text, *, fill=None, bold=False, italic=False):
        """A 列にラベル、B〜G に本文 (長文は複数行に分割)。"""
        width = self.body_width()
        chunks = chunk_text(text, width)
        for i, ch in enumerate(chunks):
            r = self.row
            self.put(r, 1, 1, label if i == 0 else "", bold=True, fill=BAND_FILL, size=9, color=INK_SUB)
            self.put(r, 2, self.ncols, ch, fill=fill, bold=bold, italic=italic)
            self.height(r, count_lines(ch, width))
            self.row += 1

    def table(self, header, rows, widths=None, label=""):
        n = len(header)
        nb = len(self.body_ws)
        # 本文列 (B〜G) を表の列数に割り当てる
        if n <= nb:
            if widths:
                total = sum(widths)
                spans, acc = [], 0
                for i, w in enumerate(widths):
                    k = max(1, round(w / total * nb)) if i < n - 1 else nb - acc
                    k = min(k, nb - acc - (n - 1 - i))
                    spans.append(k)
                    acc += k
            else:
                base = nb // n
                spans = [base] * n
                spans[-1] += nb - base * n
        else:
            spans = [1] * nb  # 列が多すぎる場合は先頭 nb 列のみ (原稿側で避ける)
            header = header[:nb]
        starts, c = [], 2
        for k in spans:
            starts.append((c, c + k - 1))
            c += k

        def write_row(values, *, head=False, band=False):
            r = self.row
            self.put(r, 1, 1, label if (head and label) else "", bold=True, fill=BAND_FILL, size=9, color=INK_SUB)
            lines = 1
            for (c1, c2), v in zip(starts, values):
                w = sum(self.body_ws[c1 - 2:c2 - 1])
                self.put(r, c1, c2, v, bold=head, fill=SUB_FILL if head else (BAND_FILL if band else None),
                         size=10 if not head else 10)
                lines = max(lines, count_lines(v, w))
            self.height(r, lines)
            self.row += 1

        write_row([str(h) for h in header], head=True)
        for i, row in enumerate(rows):
            vals = [str(x) for x in row] + [""] * (len(header) - len(row))
            write_row(vals[:len(header)], band=(i % 2 == 1))

    def code(self, text, caption=""):
        width = self.body_width()
        if caption:
            self.labeled("", caption, bold=True)
        for i, ch in enumerate(chunk_text(text, width, max_lines=26)):
            r = self.row
            self.put(r, 1, 1, "コード" if i == 0 else "", bold=True, fill=BAND_FILL, size=9, color=INK_SUB)
            self.put(r, 2, self.ncols, ch, fill=CODE_FILL, size=9)
            self.ws.row_dimensions[r].height = min(MAX_ROW_PT, count_lines(ch, width * 1.1) * 13.5 + PAD_PT)
            self.row += 1

    def blank(self, n=1):
        self.row += n

    def blocks(self, blocks):
        for kind, data in blocks:
            if kind == "p":
                self.labeled("", data)
            elif kind == "bullets":
                self.labeled("", "\n".join("・" + x for x in data))
            elif kind == "table":
                self.table(data["header"], data["rows"], data.get("widths"))
            elif kind == "code":
                self.code(data["text"], data["caption"])
            elif kind == "kid":
                self.labeled("小学生向け", data, fill=KID_FILL)
            elif kind == "note":
                self.labeled("注意・補足", data, fill=BAND_FILL)
            elif kind == "diagram":
                if data["caption"]:
                    self.labeled("", data["caption"], bold=True)
                width = self.body_width()
                for i, ch in enumerate(chunk_text(data["text"], width, max_lines=26)):
                    r = self.row
                    self.put(r, 1, 1, "図" if i == 0 else "", bold=True, fill=BAND_FILL, size=9, color=INK_SUB)
                    self.put(r, 2, self.ncols, ch, fill=WHITE, size=9)
                    self.ws.row_dimensions[r].height = min(MAX_ROW_PT, count_lines(ch, width * 1.1) * 13.5 + PAD_PT)
                    self.row += 1

    def sections(self, sections):
        for sec in sections:
            self.heading(sec["title"])
            self.blocks(sec["blocks"])
            self.blank()


def new_workbook():
    wb = Workbook()
    wb.remove(wb.active)
    # 既定フォント (Normal スタイル) も Meiryo UI にする: 空セル・列幅の基準も Meiryo UI になる
    try:
        normal = wb._named_styles["Normal"]
        normal.font = Font(name=FONT, size=10)
    except Exception:
        pass
    try:
        wb._fonts[0] = Font(name=FONT, size=10)  # 既定フォント (スタイル無しのセル)
    except Exception:
        pass
    return wb


def strip_tags(s):
    return re.sub(r"<[^>]+>", "", s)
